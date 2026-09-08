#!/usr/bin/env python3
"""
FINANSZÍROZÁSI TÖRZSEK BEOLVASÁSA — a hivatalos forrásfájlokból.

    python3 tools/ingest/torzs_ingest.py --src <könyvtár> [--out registry/kodok/tablak]

A modul terve szerint az útvonal mindig ugyanaz:

    letöltés → ingest → beágyazott JSON → build

A FORRÁSFÁJLOK NEM KERÜLNEK A REPÓBA. Hivatalos, nyilvános törzsek (NEAK,
jogszabályi mellékletek), amelyek havonta-évente frissülnek; a repó a BELŐLÜK
SZÁRMAZTATOTT táblát tartalmazza, a forrás nevével, méretével, SHA-256
lenyomatával és a kinyerés dátumával együtt (`manifest.json`). Enélkül egy
retrospektív elszámolás nem reprodukálható: nem derülne ki, MELYIK kiadásból
dolgozott a rendszer.

Amit beolvas:

  HBCS50.DBF / Hbcs50_torzs_*.xls
                                HBCS 5.0 csoportok súlyszámmal, határnapokkal
                                és közérthető megnevezéssel
  teljes_szabalykonyv_*.docx    9/2012. (II. 28.) NEFMI rendelet:
                                  2. melléklet — OENO × szakma kompetencia
                                  3. melléklet — eljárásonkénti elszámolási
                                                 szabályok (kizárások!)
                                  4. melléklet — BNO mellett elszámolható OENO
                                  5. melléklet — OENO mellett elszámolható OENO
                                  6. melléklet — telemedicinában elszámolható
  torzslista_*.xlsx             OENO törzslista PONTÉRTÉKKEL
  kompetencia_*.xlsx            OENO × szakma kompetencia (gépi olvasatú)
  BNOTORZS.DBF / BNOX_*.DBF     BNO-10 teljes magyar törzs (az újabb kiadás
                                nemet, életkort és érvényességet is ad)
  MUTET_AP.DBF                  FEKVŐBETEG beavatkozási (műtéti) törzs
  FNOTORZS.DBF                  FNO — a funkcióképesség osztályozása (ICF)
  SnomedINTL_GPSRelease_*.zip   SNOMED CT Global Patient Set (HELYI telepítés!)
  9_jegyu_beutalasi_torzslista_*.xlsx   9 jegyű finanszírozási (GYFKOD) törzs
  Diagnosztikai_torzs_*.xls     diagnosztikai anyagok árjegyzéke
  MANYAGUJ_*.xlsx               egyszerhasználatos eszközök árjegyzéke

A SNOMED KÜLÖN ÚTON JÁR — LICENCOKBÓL.

A Global Patient Set a CREATIVE COMMONS BY-ND 4.0 alatt jelenik meg:
másolni és VÁLTOZATLANUL továbbadni szabad, forrásmegjelöléssel — de az
ÁTALAKÍTOTT (adaptált) változat terjesztése NEM engedélyezett. A mi JSON
táblánk átalakítás. Ezért:

  · a SNOMED tábla a `registry/kodok/helyi/` könyvtárba készül, ami a
    `.gitignore`-ban van: HELYBEN települ, és soha nem kerül a repóba;
  · a repó legfeljebb SNOMED-AZONOSÍTÓKAT tartalmaz kereszthivatkozásként
    (mint a LOINC-kódokat), a fogalomkészletet nem;
  · a forrásmegjelölés minden SNOMED-megnevezést tartalmazó kimeneten
    kötelező — ezt a `core/coding/snomed.ts` kényszeríti ki.

AMI NINCS BENNE, ÉS EZÉRT A RENDSZER SEM ADJA:

  · a HBCS BESOROLÓ TÁBLÁZAT. A súlyszámtábla megvan, és a 10/2012. (II. 28.)
    NEFMI rendelet KÓDOLÁSI szabályai is (mit kötelező rögzíteni egy
    szüléshez); az viszont, hogy egy adott diagnózis-beavatkozás kombináció
    melyik csoportba esik — a rendelet 2. melléklete szerinti besorolási
    táblázat — nincs gépi alakban. A rendszer ezért csoportot NEM állapít meg.
  · a SNOMED CT magyar megnevezései. A fogalomkészlet betölthető (GPS), a
    validált magyar fordítás nem létezik.

AMI KORÁBBAN HIÁNYZOTT, ÉS MOST MEGVAN:

  · a FEKVŐBETEG beavatkozási törzs (MUTET_AP.DBF). NEM ICD-9-CM alakú: a
    hazai kód ÖTJEGYŰ (57410 = császármetszés, cervicalis, transversalis).
    A „74.10" alak, amit a beavatkozási törzs korábban használt, az amerikai
    ICD-9-CM-ből való, és a magyar finanszírozásban nem értelmezhető.
  · a NEM, ÉLETKOR és ÉRVÉNYESSÉG a BNO és a beavatkozási törzs sorain.
"""
import argparse, hashlib, json, os, re, sys, zipfile
from datetime import date

TXT = re.compile(r"<w:t(?: [^>]*)?>(.*?)</w:t>", re.S)
PARA = re.compile(r"<w:p[ >].*?</w:p>", re.S)
ROW = re.compile(r"<w:tr[ >].*?</w:tr>", re.S)
CELL = re.compile(r"<w:tc>(.*?)</w:tc>", re.S)
TBL = re.compile(r"<w:tbl>.*?</w:tbl>", re.S)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def find(src, pattern):
    rx = re.compile(pattern, re.I)
    for name in sorted(os.listdir(src)):
        if rx.search(name):
            return os.path.join(src, name)
    return None


def unescape(s):
    return (s.replace("&amp;", "&").replace("&lt;", "<")
             .replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'"))


def docx_paras(path):
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf8")
    return [unescape("".join(TXT.findall(p))).strip() for p in PARA.findall(xml)], xml


def docx_tables(xml):
    out = []
    for t in TBL.findall(xml):
        rows = []
        for tr in ROW.findall(t):
            rows.append([unescape("".join(TXT.findall(c))).strip() for c in CELL.findall(tr)])
        out.append([r for r in rows if any(r)])
    return out


def table(tid, label, version, valid_from, source, coverage, rows, note=None):
    t = {
        "id": tid, "label": {"hu": label}, "version": version, "validFrom": valid_from,
        "coverage": coverage, "source": source, "rows": rows,
    }
    if note:
        t["coverageNote"] = {"hu": note}
    return t


# ─── HBCS ────────────────────────────────────────────────────────────────────
def ingest_hbcs(path, tid="tbl.hbcs", version=None, valid_from=None,
                valid_until=None, cite=None):
    """
    HBCS súlyszámtábla az XLS kiadásból.

    A KORÁBBI KIADÁSOK IS BETÖLTHETŐK, és nem múzeumi célból: a 671A
    „Császármetszés" súlyszáma 2025 májusában 1,53177 volt, 2026 áprilisában
    1,74947. Egy 2025-ös eset visszamenőleges elszámolása a MAI táblával
    rossz összeget ad — csendben, mert a csoportkód ugyanaz maradt.
    """
    import xlrd
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    head = [str(sh.cell_value(0, c)).strip() for c in range(sh.ncols)]
    col = {h: i for i, h in enumerate(head)}
    ci = lambda *names: next((col[n] for n in names if n in col), None)
    c_kiem, c_foc = ci("Kiem. HBCs"), ci("Főcso-port")
    c_kod, c_min = ci("HBCs kód"), ci("Alsó határnap")
    c_max, c_norm = ci("Felső határnap"), ci("Normatív nap")
    c_suly, c_msuly = ci("Súlyszám"), ci("Műtéti súlyszám")
    c_nev, c_laik = ci("HBCs megnevezése"), ci("HBCs laikus megnevezése")
    c_hat = ci("Hatályosság")
    txt = lambda r, c: "" if c is None else str(sh.cell_value(r, c)).strip()
    rows = {}
    for r in range(1, sh.nrows):
        code = txt(r, c_kod)
        if not code:
            continue
        foc = txt(r, c_foc)
        rows[code] = {
            "label": txt(r, c_nev),
            "laikNev": txt(r, c_laik),
            "focsoport": foc,
            "focsoportSzam": foc[:2],
            "intezetiKorhozKotott": "igen" if txt(r, c_kiem) == "*" else "",
            "alsoHatarnap": txt(r, c_min),
            "felsoHatarnap": txt(r, c_max),
            "normativNap": txt(r, c_norm),
            "sulyszam": txt(r, c_suly),
            "mutetiSulyszam": txt(r, c_msuly),
            "jelzes": foc[-1] if foc[-1:] in ("P", "M", "S") else "",
            "sulyosTarsult": "igen" if code[-1:] == "Z" else "",
        }
        if c_hat is not None:
            rows[code]["hatalyossag"] = txt(r, c_hat)
    version = version or re.sub(r"[^0-9.]", "", sh.name.split("Érv.")[-1]).strip(".")
    t = table(tid, "HBCS 5.0 — homogén betegségcsoportok súlyszámmal",
              version, valid_from or "2026-04-01",
              {"cite": cite or "HBCs 5.0 törzs (NEAK közlemény)", "standard": "HBCS 5.0"},
              "full", rows, HBCS_NOTE)
    if valid_until:
        t["validUntil"] = valid_until
    return t


HBCS_NOTE = (
    "A TELJES súlyszámtábla, a csoportok KÖZÉRTHETŐ megnevezésével (laikNev) "
    "és a 10/2012. (II. 28.) NEFMI rendelet 2. § l) pontja szerinti "
    "minősítéssel (P/M/S, illetve a kód végi Z). Ami NINCS meg: a besorolási "
    "táblázat — az, hogy melyik diagnózis-beavatkozás kombináció melyik "
    "csoportba esik. A rendszer ezért csoportot NEM állapít meg; a súlyszámot "
    "és a határnapokat egy MÁR ISMERT csoporthoz adja meg.")


def ingest_hbcs_dbf(path):
    """
    HBCS 5.0 DBF kiadás — ugyanaz a súlyszámtábla, KÉT TÖBBLETTEL:

      · LAIK_NEV — a csoport KÖZÉRTHETŐ megnevezése. Ez nem díszítés: a
        zárójelentésre és a betegtájékoztatóra ez kerülhet, míg a HBCS_NEV
        („Speciális intracranialis műtétek 18 év felett") a szakmai név.
      · JEL — a `*` a MEGHATÁROZOTT INTÉZETI KÖRBEN végezhető csoportokat
        jelöli. Ha az intézet nem jogosult rá, az esetet a rendelet 20. §
        (3) bekezdése szerint a következő elszámolható legmagasabb súlyszámú
        csoportba kell sorolni — vagyis a `*` NEM címke, hanem kapu.

    A besorolási táblázat TOVÁBBRA SINCS meg: a tábla megmondja, mennyit ér egy
    csoport, azt nem, hogy melyik esetet melyik csoportba kell sorolni.
    """
    rows = {}
    for r in read_dbf(path):
        code = r["HBCS_KOD"].strip()
        if not code:
            continue
        foc = r["FOCSOP"].strip()
        rows[code] = {
            "label": r["HBCS_NEV"].strip(),
            "laikNev": r["LAIK_NEV"].strip(),
            "focsoport": foc,
            "focsoportSzam": foc[:2],
            "intezetiKorhozKotott": "igen" if r["JEL"].strip() == "*" else "",
            "alsoHatarnap": r["MIN"].strip(),
            "felsoHatarnap": r["MAX"].strip(),
            "normativNap": r["NORM"].strip(),
            "sulyszam": r["SULYSZAM"].strip(),
            "mutetiSulyszam": r["MUTSULY"].strip(),
            "jelzes": foc[-1] if foc[-1:] in ("P", "M", "S") else "",
            "sulyosTarsult": "igen" if code[-1:] == "Z" else "",
        }
    return table(
        "tbl.hbcs", "HBCS 5.0 — homogén betegségcsoportok súlyszámmal",
        "5.0/2026-04-01", "2026-04-01",
        {"cite": "HBCs 5.0 törzs (HBCS50.DBF), érvényes 2026.04.01-től (NEAK)",
         "standard": "HBCS 5.0"},
        "full", rows, HBCS_NOTE)


# ─── OENO ────────────────────────────────────────────────────────────────────
def ingest_oeno(paras, tables):
    """
    2. melléklet: OENO × szakma kompetencia.

    A tábla LAPTÖRÉSSEL KETTÉVÁGVA szerepel: a bal fele (OENO kód, Cs,
    megnevezés + az első szakmakódok) az egyik `w:tbl`, a jobb fele (a többi
    szakmakód, kód nélkül) a KÖVETKEZŐ. A két fél sorai egy az egyben
    megfelelnek egymásnak — a sorszám alapján fűzzük össze őket.
    """
    rows = {}
    for i, t in enumerate(tables):
        if not t or not t[0] or t[0][0] != "OENO kód":
            continue
        left_head = t[0][3:]
        right = tables[i + 1] if i + 1 < len(tables) else []
        right_head = right[0] if right and re.match(r"^\d+$", right[0][0] or "") else []
        specialties = left_head + right_head
        for j, r in enumerate(t[1:], start=1):
            if len(r) < 3 or not re.match(r"^\d{4,5}$", r[0]):
                continue
            cols = r[3:] + (right[j] if j < len(right) else [])
            # A lapokra tört tábla fejlécei ismétlődnek; csak a VALÓDI
            # szakmakódok (2-4 számjegy) maradnak, a visszatérő oszlopcímek nem.
            comp = [specialties[k] for k, v in enumerate(cols)
                    if k < len(specialties) and v.strip()
                    and re.match(r"^\d{2,4}$", specialties[k])]
            rows[r[0]] = {"label": r[2], "cs": r[1], "szakma": ";".join(comp)}
    return rows


def ingest_oeno_rules(paras):
    """3. melléklet: eljárásonkénti elszámolási szabályok."""
    # A 3. melléklet ÖNÁLLÓ fájlként is megjelenik (a „szabalykonyv" a
    # „teljes_szabalykonyv" nélküli kiadása). Ilyenkor nincsenek
    # melléklet-határok — az egész dokumentum a melléklet.
    i3 = next((i for i, t in enumerate(paras) if t.startswith("3. melléklet")), None)
    if i3 is None:
        seg = paras
    else:
        i4 = next((i for i, t in enumerate(paras) if t.startswith("4. melléklet")), len(paras))
        seg = paras[i3:i4]
    rules, cur = {}, None
    for t in seg:
        m = re.match(r"^(\d{4,5})\s*(\*?)\s*(.+)$", t)
        if m and re.match(r"^[A-ZÁÉÍÓÖŐÚÜŰ0-9]", m.group(3)):
            cur = m.group(1)
            rules.setdefault(cur, {"label": m.group(3).strip()})
            continue
        if cur is None:
            continue
        for label, key in (("Kizárva", "kizarva"),
                           ("Együtt 7 munkanapon belül", "egyutt7"),
                           ("Együtt", "egyutt"),
                           ("Elszámolási lehetőség (maximum)", "maxElszamolas"),
                           ("Elszámolási lehetőség", "maxElszamolas")):
            if t.startswith(label + ":"):
                val = t[len(label) + 1:].strip()
                if key in ("kizarva", "egyutt7", "egyutt"):
                    val = ";".join(c.strip() for c in re.findall(r"\d{4,5}", val))
                if val:
                    rules[cur][key] = val
                break
    return rules


def ingest_bno_conditioned(tables):
    """
    4. melléklet: „Meghatározott BNO kódok mellett elszámolható eljárások".

    Kétoszlopos tábla: egy eljárás-fejléc sort (`23.` | `46040 Első
    trimeszteri gondozási vizit…`) a hozzá tartozó BNO-sorok követik.
    """
    out, cur = {}, None
    for t in tables:
        if not t or len(t[0]) != 2:
            continue
        looks = any(re.match(r"^\d{4,5}\s*\*?\s*\D", (r[1] or "")) for r in t[:6])
        if not looks:
            continue
        for r in t:
            a, b = (r + ["", ""])[:2]
            m = re.match(r"^(\d{4,5})\s*\*?\s*(.+)$", b or "")
            if re.match(r"^\d+\.$", a or "") and m:
                cur = m.group(1)
                out.setdefault(cur, {"label": m.group(2).strip(), "bno": []})
                continue
            if cur and re.match(r"^[A-Z]\d{2,4}[A-Z0-9]*$", a or ""):
                if a not in out[cur]["bno"]:      # a melléklet ismétel kódokat
                    out[cur]["bno"].append(a)
    for k in out:
        out[k]["bno"] = ";".join(out[k]["bno"])
    return {k: v for k, v in out.items() if v["bno"]}


def ingest_flat_oeno_list(tables, header_b, drop_first=1):
    """5. és 6. melléklet: sima OENO-lista (kód | megnevezés)."""
    for t in tables:
        if len(t) < 5:
            continue
        hdr = [c for r in t[:3] for c in r]
        if header_b in hdr:
            out = {}
            for r in t:
                a = [c for c in r if c]
                if len(a) >= 3 and re.match(r"^\d{4,5}$", a[1]):
                    out[a[1]] = {"label": a[2]}
                elif len(a) == 2 and re.match(r"^\d{4,5}$", a[0]):
                    out[a[0]] = {"label": a[1]}
            if out:
                return out
    return {}


# ─── egyszerű árjegyzékek ────────────────────────────────────────────────────
def ingest_price_xls(path, tid, label, version, valid_from, cite):
    import xlrd
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    rows = {}
    for r in range(1, sh.nrows):
        v = [str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)]
        if not v[0]:
            continue
        rows[v[0]] = {"label": v[4], "csoport": v[2], "kiszereles": v[5],
                      "netto": v[6], "me": v[7]}
    return table(tid, label, version, valid_from, {"cite": cite}, "full", rows)


def ingest_price_xlsx(path, tid, label, version, valid_from, cite):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = wb[wb.sheetnames[0]]
    rows = {}
    for i, r in enumerate(sh.iter_rows(values_only=True)):
        if i == 0 or not r or not r[0]:
            continue
        v = ["" if x is None else str(x).strip() for x in r]
        rows[v[0]] = {"label": v[4], "csoport": v[2], "kiszereles": v[7],
                      "netto": v[8], "brutto": v[9], "gyarto": v[10]}
    wb.close()
    return table(tid, label, version, valid_from, {"cite": cite}, "full", rows)


# ─── DBF (dBASE III) ─────────────────────────────────────────────────────────
def read_dbf(path, enc="cp852"):
    """
    Minimális dBASE III olvasó — függőség nélkül.

    A kódolás CP852 (DOS Latin-2): a magyar egészségügyi törzsek ebben
    készültek, és rossz kódlappal az ékezetek némán elromlanak.
    """
    import struct
    with open(path, "rb") as f:
        head = f.read(32)
        n, hlen, rlen = struct.unpack("<IHH", head[4:12])
        fields = []
        while True:
            fd = f.read(32)
            if fd[0:1] in (b"\r", b""):
                break
            fields.append((fd[:11].split(b"\x00")[0].decode("latin1"), fd[16]))
        f.seek(hlen)
        out = []
        for _ in range(n):
            rec = f.read(rlen)
            if not rec or rec[0:1] == b"*":          # törölt rekord
                continue
            pos, row = 1, {}
            for name, ln in fields:
                row[name] = rec[pos:pos + ln].decode(enc, "replace").strip()
                pos += ln
            out.append(row)
    return out


def ingest_bno(path):
    """
    BNO-10 magyar törzs.

    KÉT KIADÁST ismer, és a MEZŐK alapján dönt, nem a fájlnévből: a régi
    BNOX_3_4 csak kódot, jelet és megnevezést tartalmaz, a 2025-ös BNOTORZS
    ezen felül NEMET, ÉLETKORI TARTOMÁNYT és ÉRVÉNYESSÉGET is. Ez utóbbi
    három teszi a kódot ELLENŐRIZHETŐVÉ: egy férfira kódolt O-kód nem
    stílushiba, hanem elszámolhatatlan tétel.
    """
    recs = read_dbf(path)
    rich = bool(recs) and "NEM" in recs[0]
    rows = {}
    for r in recs:
        code = r["KOD10"].strip()
        if not code:
            continue
        row = {"label": r["NEV"].strip()}
        if r.get("JEL", "").strip():
            # A kereszt (+) az alapbetegség, a csillag (*) a megnyilvánulás
            # kódja: a kettő PÁRBAN jelentendő, és külön-külön hiányos.
            row["jel"] = r["JEL"].strip()
        if rich:
            row["nem"] = r.get("NEM", "").strip()
            row["korAlso"] = r.get("KOR_A", "").strip()
            row["korFelso"] = r.get("KOR_F", "").strip()
            row["ervKezd"] = r.get("ERV_KEZD", "").strip()
            row["ervVege"] = r.get("ERV_VEGE", "").strip()
        rows[code] = row
    version, valid_from, cite = (
        ("BNOTORZS 2025-01", "1995-01-01",
         "BNO-10 magyar törzs (BNOTORZS.DBF, 2024.12.18-i kiadás)")
        if rich else
        ("BNOX 3.4", "2006-12-11", "BNO-10 magyar törzs (BNOX_3_4.DBF)"))
    note = (
        "A KÓDOK PONT NÉLKÜL szerepelnek, ahogy a hivatalos törzs tárolja "
        "őket (O141, nem O14.1). A jel mező a kereszt-csillag párokat jelöli: "
        "a + az alapbetegség, a * a megnyilvánulás kódja, és a kettő PÁRBAN "
        "jelentendő.")
    if rich:
        note += (
            " Ez a kiadás a NEMET (0 = mindkettő, 1 = férfi, 2 = nő), az "
            "ÉLETKORI TARTOMÁNYT és az ÉRVÉNYESSÉGET is tartalmazza — ezekkel "
            "a kódajánlás ellenőrizhető, nem csak kereshető.")
    return table(
        "tbl.bno", "BNO-10 — a betegségek nemzetközi osztályozása, magyar törzs",
        version, valid_from, {"cite": cite, "standard": "BNO-10"}, "full", rows, note)


def ingest_hbcs_dbf(path):
    """
    HBCS 5.0 DBF kiadás — ugyanaz a súlyszámtábla, KÉT TÖBBLETTEL:

      · LAIK_NEV — a csoport KÖZÉRTHETŐ megnevezése. Ez nem díszítés: a
        zárójelentésre és a betegtájékoztatóra ez kerülhet, míg a HBCS_NEV
        („Speciális intracranialis műtétek 18 év felett") a szakmai név.
      · JEL — a `*` a MEGHATÁROZOTT INTÉZETI KÖRBEN végezhető csoportokat
        jelöli. Ha az intézet nem jogosult rá, az esetet a rendelet 20. §
        (3) bekezdése szerint a következő elszámolható legmagasabb súlyszámú
        csoportba kell sorolni — vagyis a `*` NEM címke, hanem kapu.

    A besoroló algoritmus TOVÁBBRA SINCS meg: a tábla megmondja, mennyit ér egy
    csoport, azt nem, hogy melyik esetet melyik csoportba kell sorolni.
    """
    rows = {}
    for r in read_dbf(path):
        code = r["HBCS_KOD"].strip()
        if not code:
            continue
        rows[code] = {
            "label": r["HBCS_NEV"].strip(),
            "laikNev": r["LAIK_NEV"].strip(),
            "focsoport": r["FOCSOP"].strip(),
            "intezetiKorhozKotott": "igen" if r["JEL"].strip() == "*" else "",
            "focsoportSzam": r["FOCSOP"].strip()[:2],
            "alsoHatarnap": r["MIN"].strip(),
            "felsoHatarnap": r["MAX"].strip(),
            "normativNap": r["NORM"].strip(),
            "sulyszam": r["SULYSZAM"].strip(),
            "mutetiSulyszam": r["MUTSULY"].strip(),
            # A HBCS-kód UTOLSÓ jegye és a főcsoport jele minősít
            # (10/2012 NEFMI 2. § l) pont): P = beavatkozással jelzett,
            # M = nem jelzett, S = sürgősségi, Z (a kód végén) = súlyos
            # társult betegség. A minősítést NEM mi találjuk ki: kiolvasható.
            "jelzes": (r["FOCSOP"].strip()[-1] if r["FOCSOP"].strip()[-1:] in ("P", "M", "S") else ""),
            "sulyosTarsult": "igen" if code[-1:] == "Z" else "",
        }
    return table(
        "tbl.hbcs", "HBCS 5.0 — homogén betegségcsoportok súlyszámmal",
        "5.0/2026-04-01", "2026-04-01",
        {"cite": "HBCs 5.0 törzs (HBCS50.DBF), érvényes 2026.04.01-től (NEAK)",
         "standard": "HBCS 5.0"},
        "full", rows,
        "A TELJES súlyszámtábla, a csoportok KÖZÉRTHETŐ megnevezésével "
        "(laikNev) és a 10/2012. (II. 28.) NEFMI rendelet 2. § l) pontja "
        "szerinti minősítéssel (P/M/S, illetve a kód végi Z). Ami NINCS meg: a "
        "besoroló algoritmus — az, hogy melyik diagnózis-beavatkozás "
        "kombináció melyik csoportba esik. A rendszer ezért csoportot NEM "
        "állapít meg; a súlyszámot és a határnapokat egy MÁR ISMERT "
        "csoporthoz adja meg.")


def ingest_mut(path):
    """
    FEKVŐBETEG BEAVATKOZÁSI (MŰTÉTI) TÖRZS — a hiányzó lista.

    Ez az a törzs, amit a 24. modul `EXPECTED_STEMS` listája hiányzóként
    nevezett meg. NEM ICD-9-CM alakú: a hazai fekvőbeteg-beavatkozás ÖTJEGYŰ,
    ugyanabban a formában, mint a járóbeteg-OENO (57410 = császármetszés,
    cervicalis, transversalis) — a „74.10" alak, amit korábban használtunk,
    az amerikai ICD-9-CM-ből való, és a magyar finanszírozásban ÉRTELMEZHETETLEN.

    Három mező teszi ellenőrizhetővé a kódot, nem csak kereshetővé:

      NEM        0 = mindkét nem · 1 = férfi · 2 = nő
      ALSO/FELSO életkori tartomány ÉVBEN
      ERV_KEZD/ERV_VEGE érvényesség — egy retrospektív eset a MAI törzs
                 szerint más lehet, ezért az érvényesség az adathoz tartozik

    Egy kód ÚJRAHASZNOSÍTHATÓ: az 56900 2002-ig „Curettage uteri" volt,
    2003-tól „Terhesség-megszakítás nem orvosi indikációra". A hatályos sor
    mellé ezért a korábbi jelentés is bekerül.

    A FLAG (T/F) azt jelöli, hogy a beavatkozás önálló műtétnek számít-e.
    A jelentése a rendelethez tartozik, nem hozzánk: a mezőt ÁTVESSZÜK, de
    besorolási következtetést nem vonunk le belőle.
    """
    by_code = {}
    for r in read_dbf(path):
        code = r["KOD"].strip()
        if code:
            by_code.setdefault(code, []).append(r)
    rows = {}
    for code, recs in by_code.items():
        # EGY KÓD KÉTSZER IS SZEREPELHET, MÁS JELENTÉSSEL: az 56900 2002-ig
        # „Curettage uteri" volt, 2003-tól „Terhesség-megszakítás nem orvosi
        # indikációra". A hatályos sor a mérvadó — de a korábbi jelentést
        # MEGTARTJUK, mert egy retrospektív eset kódját enélkül némán
        # félreolvasnánk.
        recs.sort(key=lambda x: x.get("ERV_KEZD", ""))
        cur, prev = recs[-1], recs[:-1]
        rows[code] = {
            "label": cur["NEV"].strip(),
            "nem": cur.get("NEM", "").strip(),
            "korAlso": cur.get("ALSO", "").strip(),
            "korFelso": cur.get("FELSO", "").strip(),
            "flag": cur.get("FLAG", "").strip(),
            "ervKezd": cur.get("ERV_KEZD", "").strip(),
            "ervVege": cur.get("ERV_VEGE", "").strip(),
        }
        if prev:
            rows[code]["korabbiLabel"] = prev[-1]["NEV"].strip()
            rows[code]["korabbiErvVege"] = prev[-1].get("ERV_VEGE", "").strip()
    return table(
        "tbl.mut", "Fekvőbeteg beavatkozási (műtéti) törzs",
        "MUT 2026-02", "2026-02-01",
        {"cite": "NEAK fekvőbeteg beavatkozási törzs (MUTET_AP.DBF, 2026. február)",
         "standard": "OENO / fekvőbeteg beavatkozás"},
        "full", rows,
        "ÖTJEGYŰ hazai kód, nem ICD-9-CM. A nem, az életkori tartomány és az "
        "érvényesség a soron van: ezekkel a kódajánlás ELLENŐRIZHETŐ, nem csak "
        "kereshető. A FLAG mezőt átvesszük, de besorolási következtetést nem "
        "vonunk le belőle — az a besoroló algoritmus dolga, ami nincs meg.")


def ingest_fno(path):
    rows = {}
    for r in read_dbf(path):
        code = r["KOD"].strip()
        if not code:
            continue
        rows[code] = {"label": r["NEV"].strip(), "besorolas": r.get("BES_SZEMP", "").strip(),
                      "ervKezd": r.get("ERV_KEZD", "").strip(),
                      "ervVege": r.get("ERV_VEGE", "").strip()}
    return table(
        "tbl.fno", "FNO — a funkcióképesség, fogyatékosság és egészség "
                   "nemzetközi osztályozása (ICF)",
        "FNOTORZS 2019-10-03", "2004-01-01",
        {"cite": "FNO törzs (FNOTORZS.DBF); WHO ICF magyar kiadás", "standard": "FNO / ICF"},
        "full", rows,
        "A SÚLYOSSÁGI FOKOZAT A KÓD RÉSZE: a b1100 nincs problémát, a b1103 "
        "súlyos problémát jelent UGYANARRA a funkcióra. A kód utolsó jegye "
        "tehát nem alkategória, hanem MÉRTÉK — összevonva értelmetlen.")


def ingest_oeno_points(path):
    """OENO törzslista PONTÉRTÉKKEL — ez teszi elszámolhatóvá a rangsort."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = wb[wb.sheetnames[0]]
    version = re.search(r"v=(\d+)", wb.sheetnames[0])
    rows = {}
    for i, r in enumerate(sh.iter_rows(values_only=True)):
        if i == 0 or not r or not r[0]:
            continue
        v = ["" if x is None else str(x).strip() for x in r]
        rows[v[0]] = {"label": v[3], "pont": v[4], "index": v[1], "specialisFeltetel": v[2]}
    wb.close()
    return rows, (version.group(1) if version else "ismeretlen")


def ingest_kompetencia(path):
    """OENO × szakma kompetencia, gépi olvasatú kiadásból."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = wb[wb.sheetnames[0]]
    rows, head = {}, None
    for i, r in enumerate(sh.iter_rows(values_only=True)):
        v = ["" if x is None else str(x).strip() for x in r]
        if i == 0:
            head = v
            continue
        if not v or not v[0]:
            continue
        comp = [head[k] for k in range(3, len(v)) if k < len(head) and v[k]]
        rows[v[0]] = {"label": v[2], "cs": v[1], "szakma": ";".join(comp)}
    wb.close()
    return rows


SNOMED_NOTICE = (
    "© International Health Terminology Standards Development Organisation "
    "(SNOMED International). SNOMED CT® was originally created by the College "
    "of American Pathologists. A SNOMED CT Global Patient Set a Creative "
    "Commons Attribution-NoDerivatives 4.0 International licenc alatt érhető "
    "el (https://creativecommons.org/licenses/by-nd/4.0/). "
    "Forrás: https://www.snomed.org/gps"
)


def ingest_snomed(path, out_dir):
    """
    SNOMED CT Global Patient Set — HELYI telepítés.

    A kimenet SZÁNDÉKOSAN nem a `tablak/`, hanem a `helyi/` könyvtárba megy,
    ami a `.gitignore`-ban van: a CC BY-ND licenc az ÁTALAKÍTOTT változat
    terjesztését nem engedi, márpedig a JSON-ra fordítás átalakítás.
    """
    import io, zipfile
    z = zipfile.ZipFile(path)
    name = next(n for n in z.namelist() if n.endswith(".txt") and "GPSRelease" in n)
    rows = {}
    inactive = 0
    with z.open(name) as fh:
        text = io.TextIOWrapper(fh, encoding="utf8")
        text.readline()
        for line in text:
            p = line.rstrip("\n").split("\t")
            if len(p) < 4:
                continue
            if p[1] != "1":
                inactive += 1
                continue
            rows[p[0]] = {"label": p[3], "fsn": p[2]}
    version = re.search(r"(\d{8})T", os.path.basename(path))
    t = table(
        "tbl.snomed", "SNOMED CT — Global Patient Set",
        version.group(1) if version else "ismeretlen",
        f"{version.group(1)[:4]}-{version.group(1)[4:6]}-{version.group(1)[6:]}"
        if version else "",
        {"cite": SNOMED_NOTICE, "standard": "SNOMED CT GPS"},
        "full", rows,
        f"HELYI TELEPÍTÉS, a repóba NEM kerül: a CC BY-ND 4.0 licenc az "
        f"átalakított változat terjesztését nem engedi. Csak az AKTÍV "
        f"fogalmak töltődnek be ({len(rows)}); {inactive} inaktív kimarad. A "
        f"megnevezések ANGOLUL vannak — magyar fordítás nélkül a rendszer "
        f"magyar SNOMED-megnevezést NEM ad.")
    t["distribution"] = "local-only"
    t["licenceNotice"] = SNOMED_NOTICE
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "tbl-snomed.json"), "w", encoding="utf8") as f:
        json.dump(t, f, ensure_ascii=False)
        f.write("\n")
    print(f'  tbl.snomed             {len(rows):6d} aktív fogalom  '
          f'({t["version"]}) → {out_dir} [HELYI, nem terjeszthető]')
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", default="registry/kodok/tablak")
    ap.add_argument("--out-local", default="registry/kodok/helyi",
                    help="a nem terjeszthető (licencelt) törzsek helye")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    manifest = {"ingestedAt": date.today().isoformat(), "sources": [],
                "sourcesRead": [], "tables": {}}

    def record(path):
        manifest["sources"].append({
            "file": os.path.basename(path), "bytes": os.path.getsize(path),
            "sha256": sha256(path),
        })

    def write(t):
        with open(os.path.join(a.out, t["id"].replace(".", "-") + ".json"), "w",
                  encoding="utf8") as f:
            json.dump(t, f, ensure_ascii=False, indent=1)
            f.write("\n")
        manifest["tables"][t["id"]] = len(t["rows"])
        print(f'  {t["id"]:22s} {len(t["rows"]):6d} sor  ({t["version"]})')

    p = find(a.src, r"SnomedINTL_GPSRelease.*\.zip$")
    if p:
        record(p)
        n = ingest_snomed(p, a.out_local)
        manifest["tables"]["tbl.snomed (helyi)"] = n

    # A BNO-nak KÉT kiadása lehet a forráskönyvtárban. A ÚJABB nyer, és nem
    # azért, mert újabb: mert nemet, életkort és érvényességet is tartalmaz,
    # vagyis ellenőrizhető kódot ad. Egy törzsazonosítóhoz egy tábla tartozik.
    p = find(a.src, r"BNOTORZS\.DBF$") or find(a.src, r"BNOX.*\.DBF$")
    if p:
        record(p); write(ingest_bno(p))

    p = find(a.src, r"MUTET_AP\.DBF$")
    if p:
        record(p); write(ingest_mut(p))

    p = find(a.src, r"FNOTORZS\.DBF$")
    if p:
        record(p); write(ingest_fno(p))

    p = find(a.src, r"HBCS50\.DBF$")
    if p:
        record(p); write(ingest_hbcs_dbf(p))
    else:
        p = find(a.src, r"Hbcs.*\.xls$")
        if p:
            record(p); write(ingest_hbcs(p))

    # KORÁBBI HBCS KIADÁSOK — a visszamenőleges elszámoláshoz. A csoportkód
    # ugyanaz marad, a súlyszám nem: a 671A „Császármetszés" 2025 májusában
    # 1,53177 volt, 2026 áprilisában 1,74947. A mai táblával számolt 2025-ös
    # eset CSENDBEN rossz összeget ad.
    ARCHIV = [
        (r"Hbcs50_torzs_20250101(_v2)?\.xls$", "tbl.hbcs@2025-01-01",
         "5.0/2025-01-01", "2025-01-01", "2025-04-30",
         "HBCs 5.0 törzs, érvényes 2025.01.01-től (NEAK közlemény)"),
        (r"Hbcs50_torzs_20241101_20250101_20250501\.xls$", "tbl.hbcs@2025-05-01",
         "5.0/2025-05-01", "2025-05-01", "2026-03-31",
         "HBCs 5.0 törzs, érvényes 2025.05.01-től (NEAK közlemény)"),
    ]
    for rx, tid, ver, vf, vu, cite in ARCHIV:
        p = find(a.src, rx)
        if p:
            record(p); write(ingest_hbcs(p, tid, ver, vf, vu, cite))

    p = find(a.src, r"szabalykonyv.*\.docx$|teljes_szabalykonyv")
    if p:
        record(p)
        paras, xml = docx_paras(p)
        tables = docx_tables(xml)
        src = {"cite": "9/2012. (II. 28.) NEFMI rendelet az Egészségbiztosítási Alap "
                       "terhére finanszírozható járóbeteg-szakellátási tevékenységekről "
                       "(hatályos 2026.01.01.)", "standard": "9/2012 NEFMI"}
        # A KOMPETENCIA és a PONTÉRTÉK gépi olvasatú kiadásból jön, ha van:
        # a docx táblái lapokra vannak törve, és az összefűzésük hibázhat.
        pk = find(a.src, r"^kompetencia.*\.xlsx$")
        pt = find(a.src, r"^torzslista.*\.xlsx$")
        oeno = ingest_kompetencia(pk) if pk else ingest_oeno(paras, tables)
        if pk:
            record(pk)
        points_version = None
        if pt:
            record(pt)
            pts, points_version = ingest_oeno_points(pt)
            for code, row in pts.items():
                oeno.setdefault(code, {"label": row["label"], "cs": "", "szakma": ""})
                for k in ("pont", "index", "specialisFeltetel"):
                    if row.get(k):
                        oeno[code][k] = row[k]
        rules = ingest_oeno_rules(paras)
        for code, r in rules.items():
            oeno.setdefault(code, {"label": r.get("label", ""), "cs": "", "szakma": ""})
            for k in ("kizarva", "egyutt7", "egyutt", "maxElszamolas"):
                if r.get(k):
                    oeno[code][k] = r[k]
        write(table("tbl.oeno", "OENO — járóbeteg-szakellátási eljárások",
                    "9/2012 NEFMI, hatályos 2026-01-01", "2026-01-01", src, "full", oeno,
                    "A rendelet 2. és 3. mellékletéből: eljáráskód, megnevezés, szakmai "
                    "kompetencia, kizárási és együtt-jelenthetőségi szabályok, és a "
                    "törzslista PONTÉRTÉKE. A pontérték teszi elvégezhetővé azt, amit a "
                    "jogszabály 5. § (7) előír: kizárásos ütközésnél a magasabb "
                    "pontszámú eljárás számolható el."))
        write(table("tbl.oeno.bno", "BNO mellett elszámolható eljárások (4. melléklet)",
                    "9/2012 NEFMI, hatályos 2026-01-01", "2026-01-01", src, "full",
                    ingest_bno_conditioned(tables),
                    "Az eljárás CSAK a felsorolt BNO-kódok mellett számolható el. Ez a "
                    "rendszer szempontjából a legfontosabb melléklet: a kódajánlás "
                    "diagnózis-feltételét jogszabály mondja meg, nem mi."))
        write(table("tbl.oeno.egyutt", "99985 mellett elszámolható eljárások (5. melléklet)",
                    "9/2012 NEFMI, hatályos 2026-01-01", "2026-01-01", src, "full",
                    ingest_flat_oeno_list(tables, "Tevékenység megnevezése")))
        write(table("tbl.oeno.tele", "Telemedicinában elszámolható eljárások (6. melléklet)",
                    "9/2012 NEFMI, hatályos 2026-01-01", "2026-01-01", src, "full",
                    ingest_flat_oeno_list(tables, "Egészségügyi eljárás")))

    p = find(a.src, r"9.?jegy.*\.xlsx$")
    if p:
        record(p)
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        sh = wb["teljes törzs"]
        rows = {}
        for i, r in enumerate(sh.iter_rows(values_only=True)):
            if i == 0 or not r or not r[6]:
                continue
            v = ["" if x is None else str(x).strip() for x in r]
            # A SZEMÉLYNÉV KIMARAD: a törzs orvosneveket is tartalmaz, azok a
            # kódajánláshoz nem kellenek. Ami kell: melyik intézmény melyik
            # szakmára jogosult finanszírozott ellátást jelenteni.
            rows[v[6]] = {"label": v[8], "intkod": v[3], "varmegye": v[2]}
        wb.close()
        write(table("tbl.gyfkod", "9 jegyű finanszírozási (GYFKOD) törzs",
                    str(rows and next(iter(rows)) and "202609"), "2026-09-01",
                    {"cite": "NEAK: 9 jegyű beutalási törzslista, 2026. szeptember"},
                    "full", rows,
                    "A SZEMÉLYNEVEK (GYFKOD_KAPCS_NEV orvosnevei) SZÁNDÉKOSAN KIMARADTAK: "
                    "a kódajánláshoz nem kellenek, és nyilvános törzsből származó "
                    "személyes adat fölösleges másolása."))

    p = find(a.src, r"Diagnosztikai_torzs.*\.xls$")
    if p:
        record(p)
        write(ingest_price_xls(p, "tbl.diag", "Diagnosztikai anyagok törzse",
                               "2005-12-20", "2005-12-20",
                               "Diagnosztikai törzs, 2005.12.20. (archív kiadás)"))

    p = find(a.src, r"MANYAG.*\.xlsx$")
    if p:
        record(p)
        write(ingest_price_xlsx(p, "tbl.eszkoz", "Egyszerhasználatos eszközök törzse",
                                "2018-06-04", "2018-06-04",
                                "Egyszerhasználatos (műanyag) eszköz törzs, 2018.06.04."))

    # KÉZZEL FELDOLGOZOTT FORRÁSOK. Ezekből nem tábla lesz, hanem szabály:
    # a jogszabály szövegét és a NEAK technikai útmutatót ember olvasta, és a
    # `registry/kodok/fekvo/` állományaiba írta át. A lenyomatuk mégis ide
    # tartozik: enélkül nem derülne ki, MELYIK kiadásból való a szabály.
    for rx, mit in [
        (r"^fekvo\.docx$|10_2012.*\.docx$",
         "10/2012. (II. 28.) NEFMI rendelet a HBCs kódolási és besorolási "
         "szabályairól → registry/kodok/fekvo/"),
        (r"^tech\.pdf$|Technikai_leiras.*\.pdf$",
         "NEAK: Fekvőbeteg jelentés — technikai útmutató → "
         "registry/kodok/fekvo/rekordkep.json"),
    ]:
        p = find(a.src, rx)
        if p:
            manifest["sourcesRead"].append({
                "file": os.path.basename(p), "bytes": os.path.getsize(p),
                "sha256": sha256(p), "hasznalat": mit,
            })
            print(f'  {os.path.basename(p):40s} kézzel feldolgozva')

    mpath = os.path.join(os.path.dirname(a.out.rstrip("/")) or ".", "manifest.json")
    with open(mpath, "w", encoding="utf8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print("manifest:", json.dumps(manifest["tables"], ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
